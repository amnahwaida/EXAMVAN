"""EXAMVAN routes — Submissions management, detail, export CSV/XLSX."""
import json
import csv
import io
import os
from datetime import datetime, timezone

from flask import (
    request, jsonify, render_template,
    redirect, url_for, session, flash, abort, send_file
)

from app import (
    app, get_db,
    admin_required,
    check_exam_ownership, check_submission_ownership,
    get_network_info, BASE_DIR, VERSION,
)
from helpers import (
    localize_date_string, format_iso_utc, _normalize_q_num,
    evaluate_answers_detailed, calculate_submission_score,
    parse_roles,
)
from routes._shared import (
    error_response, success_response,
    _csv_safe, _sanitize_xlsx,
    _parse_identity_fields, _get_identity_data, _get_sub_field,
    identity_headers_and_values, sanitize_sheet_name,
)

from app import logger as app_logger
logger = app_logger

# Optional openpyxl for XLSX export
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ===== Submissions Overview Page =====

@app.route('/admin/submissions')
@admin_required
def admin_submissions():
    """Submissions overview page for admin with pagination."""
    if not session.get('is_super_admin') and not session.get('is_operator') and 'guru' not in session.get('user_roles', []):
        flash('Akses ditolak. Halaman ini khusus untuk Guru.', 'error')
        return redirect(url_for('admin_pengawas'))
    db = get_db()
    is_super_admin = session.get('is_super_admin', False)
    is_operator = session.get('is_operator', False)

    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 25, type=int)
    exam_filter = request.args.get('exam_id', type=int)
    per_page = min(max(per_page, 5), 100)

    conditions = []
    params = []
    if is_super_admin:
        pass
    elif is_operator:
        op_row = db.execute('SELECT instansi FROM admin_users WHERE id = ?', (session['admin_id'],)).fetchone()
        op_instansi = op_row['instansi'] if op_row else ''
        conditions.append('(e.created_by = ? OR e.created_by IN (SELECT id FROM admin_users WHERE instansi = ?))')
        params.extend([session['admin_id'], op_instansi])
    else:
        conditions.append('(e.created_by = ? OR s.exam_id IN (SELECT exam_id FROM exam_pengawas WHERE user_id = ?))')
        params.extend([session['admin_id'], session['admin_id']])
    if exam_filter:
        conditions.append('s.exam_id = ?')
        params.append(exam_filter)

    where_clause = (' WHERE ' + ' AND '.join(conditions)) if conditions else ''

    count_sql = f'SELECT COUNT(*) as cnt FROM submissions s JOIN exams e ON s.exam_id = e.id{where_clause}'
    total_submissions = db.execute(count_sql, params).fetchone()['cnt']

    base_query = (
        'SELECT s.*, e.name as exam_name, e.questions_json '
        'FROM submissions s JOIN exams e ON s.exam_id = e.id'
    )
    order = ' ORDER BY s.created_at DESC'
    limit_offset = ' LIMIT ? OFFSET ?'
    submissions = db.execute(base_query + where_clause + order + limit_offset, params + [per_page, (page - 1) * per_page]).fetchall()

    # Compute max_score and percentage per submission
    sub_data = []
    for sub in submissions:
        sub_max_score = None
        if sub['questions_json']:
            try:
                questions = json.loads(sub['questions_json'])
                sub_max_score = sum(float(q.get('weight', 1.0)) for q in questions)
            except Exception as e:
                logger.warning("Failed to parse questions for submission %s: %s", sub['id'], e)
        try:
            identity_data = json.loads(sub['identity_data']) if sub['identity_data'] else {}
        except Exception:
            identity_data = {}
        sub_data.append({
            'id': sub['id'],
            'exam_id': sub['exam_id'],
            'exam_name': sub['exam_name'],
            'student_name': sub['student_name'],
            'exam_number': sub['exam_number'],
            'student_class': sub['student_class'],
            'identity_data': identity_data,
            'answers_json': sub['answers_json'],
            'score': sub['score'],
            'max_score': sub_max_score,
            'score_pct': round((sub['score'] / sub_max_score * 100), 1) if (sub['score'] is not None and sub_max_score) else None,
            'start_time': sub['start_time'],
            'created_at': sub['created_at'],
            'mac_address': sub['mac_address'],
        })

    if is_super_admin:
        exams = db.execute('SELECT id, name FROM exams ORDER BY created_at DESC').fetchall()
    elif is_operator:
        op_row = db.execute('SELECT instansi FROM admin_users WHERE id = ?', (session['admin_id'],)).fetchone()
        op_instansi = op_row['instansi'] if op_row else ''
        exams = db.execute('SELECT id, name FROM exams WHERE created_by = ? OR created_by IN (SELECT id FROM admin_users WHERE instansi = ?) ORDER BY created_at DESC', (session['admin_id'], op_instansi)).fetchall()
    else:
        exams = db.execute('SELECT id, name FROM exams WHERE created_by = ? OR id IN (SELECT exam_id FROM exam_pengawas WHERE user_id = ?) ORDER BY created_at DESC', (session['admin_id'], session['admin_id'])).fetchall()

    total_pages = max(1, (total_submissions + per_page - 1) // per_page)

    max_score = None
    exam_info = None
    if exam_filter:
        exam_row = db.execute('SELECT id, name, token, created_at, size_bytes, questions_json, start_time, end_time FROM exams WHERE id = ?', (exam_filter,)).fetchone()
        if exam_row:
            if exam_row['questions_json']:
                try:
                    questions = json.loads(exam_row['questions_json'])
                    max_score = sum(float(q.get('weight', 1.0)) for q in questions)
                except Exception as e:
                    logger.warning("Failed to parse questions for exam filter %s: %s", exam_filter, e)
            sub_count = db.execute('SELECT COUNT(*) as cnt FROM submissions WHERE exam_id = ?', (exam_filter,)).fetchone()['cnt']
            exam_info = {
                'name': exam_row['name'],
                'token': exam_row['token'],
                'created_at': exam_row['created_at'],
                'size_mb': round((exam_row['size_bytes'] or 0) / (1024 * 1024), 2),
                'sub_count': sub_count,
                'start_time': exam_row['start_time'],
                'end_time': exam_row['end_time'],
            }

    local_ip = get_network_info()['display_host']

    return render_template(
        'submissions.html',
        submissions=sub_data,
        exams=exams,
        local_ip=local_ip,
        admin_user=session.get('admin_username', 'Admin'),
        admin_role=session.get('admin_role', 'guru'),
        active_page='submissions',
        page=page,
        per_page=per_page,
        total_pages=total_pages,
        total_submissions=total_submissions,
        exam_filter_param=exam_filter or '',
        exam_info=exam_info,
    )


# ===== Submission Detail =====

@app.route('/admin/api/submissions/<int:submission_id>/detail')
@admin_required
def admin_submission_detail(submission_id):
    """Get detailed student answers compared with keys."""
    db = get_db()
    if not check_submission_ownership(db, submission_id):
        return error_response('Akses ditolak: Anda tidak memiliki akses ke data ini', 403)
    sub = db.execute(
        'SELECT s.*, e.name as exam_name, e.questions_json '
        'FROM submissions s JOIN exams e ON s.exam_id = e.id '
        'WHERE s.id = ?', (submission_id,)
    ).fetchone()

    if not sub:
        return error_response('Hasil ujian tidak ditemukan', 404)

    try:
        answers = json.loads(sub['answers_json'])
    except Exception:
        answers = {}

    try:
        questions = json.loads(sub['questions_json']) if sub['questions_json'] else []
    except Exception:
        questions = []

    evaluated_answers = evaluate_answers_detailed(answers, questions)

    try:
        identity_data = json.loads(sub['identity_data']) if sub['identity_data'] else {}
    except Exception:
        identity_data = {}

    return jsonify({
        'success': True,
        'submission_id': sub['id'],
        'student_name': sub['student_name'],
        'exam_number': sub['exam_number'],
        'student_class': sub['student_class'],
        'identity_data': identity_data,
        'exam_name': sub['exam_name'],
        'score': sub['score'],
        'start_time': format_iso_utc(sub['start_time']) if sub['start_time'] else None,
        'mac_address': sub['mac_address'],
        'created_at': format_iso_utc(sub['created_at']),
        'answers': answers,
        'questions': questions,
        'evaluated_answers': evaluated_answers
    })


# ===== Export Single Submission (CSV) =====

@app.route('/admin/api/submissions/<int:submission_id>/export_detail')
@admin_required
def admin_export_submission_detail(submission_id):
    """Export a single student's detailed answers to CSV."""
    db = get_db()
    if not check_submission_ownership(db, submission_id):
        return abort(403)
    sub = db.execute(
        'SELECT s.*, e.name as exam_name, e.questions_json '
        'FROM submissions s JOIN exams e ON s.exam_id = e.id '
        'WHERE s.id = ?', (submission_id,)
    ).fetchone()

    if not sub:
        return abort(404)

    try:
        answers = json.loads(sub['answers_json'])
    except Exception:
        answers = {}

    try:
        questions = json.loads(sub['questions_json']) if sub['questions_json'] else []
    except Exception:
        questions = []

    tz_offset = request.args.get('tz_offset', type=int)

    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow(['Detail Hasil Ujian Siswa'])
    cw.writerow(['Nama Ujian', _csv_safe(sub['exam_name'])])
    cw.writerow(['Nama Siswa', _csv_safe(sub['student_name'])])
    cw.writerow(['Nomor Ujian', _csv_safe(sub['exam_number'])])
    cw.writerow(['Kelas', _csv_safe(sub['student_class'])])
    cw.writerow(['Nilai Akhir', sub['score'] if sub['score'] is not None else 'Belum Dinilai'])
    cw.writerow(['Waktu Mulai', localize_date_string(sub['start_time'], tz_offset) if sub['start_time'] else '—'])
    cw.writerow(['Waktu Kumpul', localize_date_string(sub['created_at'], tz_offset)])
    cw.writerow(['ID Perangkat', _csv_safe(sub['mac_address'] or '—')])
    cw.writerow([])
    cw.writerow(['No. Soal', 'Tipe Soal', 'Bobot Maks', 'Jawaban Siswa', 'Kunci Jawaban', 'Status', 'Poin Didapat'])

    evaluated = evaluate_answers_detailed(answers, questions)

    for q in questions:
        q_num = _normalize_q_num(q.get('number', ''))
        student_ans = answers.get(q_num)
        correct_ans = q.get('key')
        q_weight = float(q.get('weight', 1.0))

        eval_info = evaluated.get(q_num, {})
        earned_q_weight = eval_info.get('earned', 0.0)
        status_text = eval_info.get('statusText', 'unanswered')

        student_ans_str = ''
        if isinstance(student_ans, list):
            student_ans_str = ', '.join(student_ans)
        elif isinstance(student_ans, dict):
            student_ans_str = ', '.join([f"{k}:{v}" for k, v in student_ans.items()])
        elif student_ans is not None:
            student_ans_str = str(student_ans)

        correct_ans_str = ''
        if isinstance(correct_ans, list):
            correct_ans_str = ', '.join(correct_ans)
        elif isinstance(correct_ans, dict):
            correct_ans_str = ', '.join([f"{k}:{v}" for k, v in correct_ans.items()])
        elif correct_ans is not None:
            correct_ans_str = str(correct_ans)

        type_labels = {
            'single_choice': 'Pilihan Ganda',
            'multiple_choice': 'PG Kompleks',
            'true_false': 'Benar / Salah',
            'matching': 'Menjodohkan',
            'short_answer': 'Isian Singkat',
        }
        cw.writerow([
            q['number'],
            type_labels.get(q['type'], q['type']),
            q_weight,
            _csv_safe(student_ans_str),
            _csv_safe(correct_ans_str),
            status_text,
            round(earned_q_weight, 2)
        ])

    output = si.getvalue()
    si.close()

    filename = f"detail_jawaban_{sub['student_name'].replace(' ', '_')}_{sub['exam_number']}.csv"
    response = send_file(
        io.BytesIO(output.encode('utf-8-sig')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=filename
    )
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response


# ===== Delete Submission =====

@app.route('/admin/api/submissions/<int:submission_id>', methods=['DELETE'])
@admin_required
def admin_delete_submission(submission_id):
    """Delete a student submission."""
    db = get_db()
    if not check_submission_ownership(db, submission_id):
        return error_response('Akses ditolak: Anda tidak memiliki akses ke data ini', 403)
    db.execute('DELETE FROM submissions WHERE id = ?', (submission_id,))
    db.commit()
    return jsonify({'success': True, 'message': 'Hasil ujian berhasil dihapus'})


# ===== Export Submissions (CSV / XLSX) =====

@app.route('/admin/api/submissions/export')
@admin_required
def admin_export_submissions():
    """Export student submissions. Generates multi-sheet XLSX for specific exam, CSV for all."""
    exam_id = request.args.get('exam_id')
    tz_offset = request.args.get('tz_offset', type=int)

    db = get_db()
    is_super_admin = session.get('is_super_admin', False)
    is_operator = session.get('is_operator', False)

    # --- Multi-sheet XLSX export for a specific exam ---
    if exam_id:
        exam = db.execute('SELECT * FROM exams WHERE id = ?', (exam_id,)).fetchone()
        if not exam:
            return error_response('Ujian tidak ditemukan', 404)
        if not is_super_admin and not is_operator and exam['created_by'] != session['admin_id']:
            return error_response('Akses ditolak', 403)

        submissions = db.execute(
            'SELECT * FROM submissions WHERE exam_id = ? ORDER BY student_class, student_name',
            (exam_id,)
        ).fetchall()

        MAX_XLSX_STUDENTS = 500
        if len(submissions) > MAX_XLSX_STUDENTS:
            return error_response(f'Export dibatasi maksimal {MAX_XLSX_STUDENTS} siswa per file. Gunakan filter untuk mengurangi jumlah.', 400)

        try:
            questions = json.loads(exam['questions_json']) if exam['questions_json'] else []
        except Exception:
            questions = []

        # Try XLSX first
        if OPENPYXL_AVAILABLE:
            try:
                return _generate_exam_xlsx(exam, submissions, questions, tz_offset)
            except ImportError:
                pass
            except Exception as e:
                logger.error("XLSX export error: %s", e)

        # CSV fallback for specific exam
        si = io.StringIO()
        cw = csv.writer(si)
        ident_fields = _parse_identity_fields(exam)
        ident_headers, _ = identity_headers_and_values(ident_fields, {}, submissions[0] if submissions else {'student_name': '', 'exam_number': '', 'student_class': ''})
        headers = ['ID', 'Nama Ujian'] + ident_headers + ['Nilai', 'Waktu Mulai', 'Waktu Kumpul', 'ID Perangkat']
        cw.writerow(headers)
        for row in submissions:
            idata = _get_identity_data(row)
            _, ident_vals = identity_headers_and_values(ident_fields, idata, row)
            cw.writerow([
                row['id'], _csv_safe(exam['name']),
                *ident_vals,
                row['score'] if row['score'] is not None else 'Belum Dinilai',
                localize_date_string(row['start_time'], tz_offset) if row['start_time'] else '—',
                localize_date_string(row['created_at'], tz_offset),
                _csv_safe(row['mac_address'] or '—')
            ])
        output = si.getvalue(); si.close()
        safe_exam_name = ''.join(c for c in exam['name'] if c.isalnum() or c in ' _-').strip().replace(' ', '_')
        filename = f"Hasil_Ujian_{safe_exam_name}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.csv"
        response = send_file(io.BytesIO(output.encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name=filename)
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
        return response

    # --- CSV export for all exams ---
    query = (
        'SELECT s.id, e.name as exam_name, s.student_name, s.exam_number, s.student_class, s.identity_data, s.score, s.start_time, s.mac_address, s.created_at '
        'FROM submissions s JOIN exams e ON s.exam_id = e.id'
    )
    conditions = []
    params = []

    if not is_super_admin and not is_operator:
        conditions.append('(e.created_by = ? OR s.exam_id IN (SELECT exam_id FROM exam_pengawas WHERE user_id = ?))')
        params.extend([session['admin_id'], session['admin_id']])

    if conditions:
        query += ' WHERE ' + ' AND '.join(conditions)

    query += ' ORDER BY s.created_at DESC'
    all_submissions = db.execute(query, params).fetchall()

    all_ident_keys = []
    all_ident_keys_set = set()
    for row in all_submissions:
        idata = _get_identity_data(row)
        for k in idata:
            if k not in all_ident_keys_set:
                all_ident_keys_set.add(k)
                all_ident_keys.append(k)

    si = io.StringIO()
    cw = csv.writer(si)
    fixed_headers = ['ID', 'Nama Ujian']
    ident_headers_csv = [k.replace('_', ' ').title() for k in all_ident_keys]
    trailing_headers = ['Nilai', 'Waktu Mulai', 'Waktu Kumpul', 'ID Perangkat']
    cw.writerow(fixed_headers + ident_headers_csv + trailing_headers)

    for row in all_submissions:
        idata = _get_identity_data(row)
        ident_vals = [_csv_safe(str(idata.get(k, '—'))) if idata.get(k) else '—' for k in all_ident_keys]
        cw.writerow([
            row['id'], _csv_safe(row['exam_name']),
            *ident_vals,
            row['score'] if row['score'] is not None else 'Belum Dinilai',
            localize_date_string(row['start_time'], tz_offset) if row['start_time'] else '—',
            localize_date_string(row['created_at'], tz_offset),
            _csv_safe(row['mac_address'] or '—')
        ])

    output = si.getvalue(); si.close()
    filename = f"hasil_ujian_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.csv"
    response = send_file(io.BytesIO(output.encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name=filename)
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response


# =====================================================================
# XLSX Export Helpers
# =====================================================================


def _set_column_widths(ws, widths):
    """Set column widths for a worksheet using an ordered list of widths."""
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w


def _excel_status_color(cell, status_text):
    """Apply color formatting to a status cell based on evaluation text."""
    if status_text == 'correct':
        cell.fill = PatternFill(start_color='d1fae5', end_color='d1fae5', fill_type='solid')
        cell.font = Font(color='065f46', bold=True)
    elif status_text == 'partial':
        cell.fill = PatternFill(start_color='fef3c7', end_color='fef3c7', fill_type='solid')
        cell.font = Font(color='92400e', bold=True)
    elif status_text == 'incorrect':
        cell.fill = PatternFill(start_color='fee2e2', end_color='fee2e2', fill_type='solid')
        cell.font = Font(color='991b1b', bold=True)
    else:
        cell.fill = PatternFill(start_color='f3f4f6', end_color='f3f4f6', fill_type='solid')
        cell.font = Font(color='374151')


def _build_question_detail_row(ws, row_num, q_num, q, student_answers, evaluation):
    """Build a single question detail row. Returns (q_weight, earned) tuple."""
    q_weight = float(q.get('weight', 1.0))
    eval_info = evaluation.get(q_num, {})
    earned = eval_info.get('earned', 0.0)
    status_text = eval_info.get('statusText', 'unanswered')

    student_ans = student_answers.get(q_num)
    correct_ans = q.get('key')

    def _fmt(val):
        if val is None:
            return '—'
        if isinstance(val, list):
            return ', '.join(str(x) for x in val)
        if isinstance(val, dict):
            return ', '.join(f'{k} ➔ {v}' for k, v in val.items())
        return str(val)

    student_display = _fmt(student_ans) if student_ans is not None and student_ans != '' else '—'
    key_display = _fmt(correct_ans)

    type_labels = {
        'single_choice': 'Pilihan Ganda',
        'multiple_choice': 'PG Kompleks',
        'true_false': 'Benar / Salah',
        'matching': 'Menjodohkan',
        'short_answer': 'Isian Singkat',
    }

    thin_border = Border(
        left=Side(style='thin', color='d1d5db'),
        right=Side(style='thin', color='d1d5db'),
        top=Side(style='thin', color='d1d5db'),
        bottom=Side(style='thin', color='d1d5db'),
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    row_values = [
        q['number'],
        type_labels.get(q['type'], q['type']),
        q_weight,
        student_display,
        key_display,
        status_text,
        round(earned, 2),
    ]
    for col_idx, val in enumerate(row_values, 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.border = thin_border
        cell.alignment = center_align if col_idx in [1, 3, 6, 7] else left_align
        cell.value = val
        if col_idx == 6:
            _excel_status_color(cell, str(val))

    return q_weight, earned


def _build_exam_summary_sheet(ws, exam, submissions, tz_offset):
    """Build Sheet 1: Ringkasan Hasil — title, metadata, and student summary table."""
    NAVY = '1e3a8a'
    GREEN_BG = 'd1fae5'
    GREEN_FG = '065f46'
    GRAY_BG = 'f3f4f6'
    GRAY_FG = '374151'

    ws.title = 'Ringkasan Hasil'
    ws.sheet_properties.tabColor = NAVY

    thin_border = Border(
        left=Side(style='thin', color='d1d5db'),
        right=Side(style='thin', color='d1d5db'),
        top=Side(style='thin', color='d1d5db'),
        bottom=Side(style='thin', color='d1d5db'),
    )
    title_font = Font(bold=True, size=14, color=NAVY)
    meta_label_font = Font(bold=True, size=11)
    meta_value_font = Font(size=11)
    header_font = Font(bold=True, color='ffffff', size=11)
    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def _style_header_row(ws, row, col_count):
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    def _style_data_cell(ws, row, col, align='center'):
        cell = ws.cell(row=row, column=col)
        cell.border = thin_border
        cell.alignment = center_align if align == 'center' else left_align
        return cell

    ident_fields = _parse_identity_fields(exam)
    ident_headers, _ = identity_headers_and_values(ident_fields, {}, {'student_name': '', 'exam_number': '', 'student_class': ''})
    num_ident = len(ident_headers)
    num_total = 1 + num_ident + 4

    title_range = f'A1:{get_column_letter(num_total)}1'
    ws.merge_cells(title_range)
    title_cell = ws['A1']
    title_cell.value = f'RINGKASAN HASIL UJIAN — {exam["name"]}'
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal='left', vertical='center')

    meta_data = [
        ('Token Ujian:', exam['token'] or '—'),
        ('Total Peserta:', str(len(submissions))),
        ('Tanggal Export:', localize_date_string(datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S'), tz_offset)),
    ]
    for i, (label, value) in enumerate(meta_data):
        ws.cell(row=3 + i, column=1, value=label).font = meta_label_font
        ws.cell(row=3 + i, column=2, value=value).font = meta_value_font

    summary_headers = ['No.'] + ident_headers + ['Nilai Akhir', 'Status', 'Waktu Mulai', 'Waktu Pengumpulan', 'ID Perangkat']
    header_row = 7
    for col_idx, h in enumerate(summary_headers, 1):
        ws.cell(row=header_row, column=col_idx, value=h)
    _style_header_row(ws, header_row, len(summary_headers))

    for i, sub in enumerate(submissions):
        row_num = header_row + 1 + i
        score = sub['score']
        score_display = round(score, 2) if score is not None else 'Belum Dinilai'
        status = 'Sudah Dinilai' if score is not None else 'Belum Dinilai'

        idata = _get_identity_data(sub)
        _, ident_vals = identity_headers_and_values(ident_fields, idata, sub)

        values = [i + 1] + ident_vals + [score_display, status,
                  localize_date_string(sub['start_time'], tz_offset) if sub['start_time'] else '—',
                  localize_date_string(sub['created_at'], tz_offset),
                  _sanitize_xlsx(sub['mac_address'] or '—')]
        for col_idx, val in enumerate(values, 1):
            cell = _style_data_cell(ws, row_num, col_idx,
                                   'center' if col_idx in [1, num_ident + 2, num_ident + 3] else 'left')
            cell.value = val
            score_col = num_ident + 2
            status_col = num_ident + 3
            if col_idx == score_col and score is not None:
                cell.font = Font(bold=True)
            if col_idx == status_col:
                if score is not None:
                    cell.fill = PatternFill(start_color=GREEN_BG, end_color=GREEN_BG, fill_type='solid')
                    cell.font = Font(color=GREEN_FG, bold=True)
                else:
                    cell.fill = PatternFill(start_color=GRAY_BG, end_color=GRAY_BG, fill_type='solid')
                    cell.font = Font(color=GRAY_FG)

    for col_idx in range(1, len(summary_headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = \
            max(14, len(summary_headers[col_idx - 1]) + 6)


def _build_student_detail_sheet(wb, sub, questions, tz_offset, used_names, ident_fields=None):
    """Build a per-student detail sheet with info header and question-by-question results."""
    NAVY = '1e3a8a'
    LIGHT_BLUE = 'dbeafe'

    idata = _get_identity_data(sub)
    idata_name = str(idata.get('student_name', sub['student_name']) or sub['student_name'])
    idata_num = str(idata.get('exam_number', sub['exam_number']) or sub['exam_number'])
    base_name = f"{idata_num} {idata_name}"
    sheet_name = sanitize_sheet_name(base_name)
    if sheet_name in used_names:
        counter = 2
        while f"{sheet_name[:28]}_{counter}" in used_names:
            counter += 1
        sheet_name = f"{sheet_name[:28]}_{counter}"
    used_names.add(sheet_name)

    ws = wb.create_sheet(title=sheet_name)

    thin_border = Border(
        left=Side(style='thin', color='d1d5db'),
        right=Side(style='thin', color='d1d5db'),
        top=Side(style='thin', color='d1d5db'),
        bottom=Side(style='thin', color='d1d5db'),
    )
    title_font = Font(bold=True, size=14, color=NAVY)
    meta_label_font = Font(bold=True, size=11)
    meta_value_font = Font(size=11)
    header_font = Font(bold=True, color='ffffff', size=11)
    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    def _style_header_row(ws, row, col_count):
        for col_idx in range(1, col_count + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

    def _style_data_cell(ws, row, col, align='center'):
        cell = ws.cell(row=row, column=col)
        cell.border = thin_border
        cell.alignment = center_align if align == 'center' else left_align
        return cell

    ws.merge_cells('A1:G1')
    ws['A1'].value = 'DETAIL HASIL UJIAN SISWA'
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    info_rows = []
    if ident_fields:
        for f in ident_fields:
            key = f['key']
            label = f['label']
            val = idata.get(key)
            if val is None or val == '':
                if key == 'student_name': val = sub['student_name']
                elif key == 'exam_number': val = sub['exam_number']
                elif key == 'student_class': val = sub['student_class']
                else: val = '—'
            info_rows.append((f'{label}:', _sanitize_xlsx(str(val)) if val is not None else '—'))
    else:
        info_rows = [
            ('Nama Siswa:', _sanitize_xlsx(sub['student_name'])),
            ('Nomor Ujian:', _sanitize_xlsx(sub['exam_number'])),
            ('Kelas:', _sanitize_xlsx(sub['student_class'])),
        ]
    info_rows += [
        ('Nilai Akhir:', round(sub['score'], 2) if sub['score'] is not None else 'Belum Dinilai'),
        ('Waktu Mulai:', localize_date_string(sub['start_time'], tz_offset) if sub['start_time'] else '—'),
        ('Waktu Pengumpulan:', localize_date_string(sub['created_at'], tz_offset)),
        ('ID Perangkat:', _sanitize_xlsx(sub['mac_address'] or '—')),
    ]
    for i, (label, value) in enumerate(info_rows):
        ws.cell(row=3 + i, column=1, value=label).font = meta_label_font
        val_cell = ws.cell(row=3 + i, column=2, value=value)
        val_cell.font = meta_value_font
        if label == 'Nilai Akhir:' and sub['score'] is not None:
            val_cell.font = Font(bold=True, size=12, color=NAVY)

    detail_header_row = 3 + len(info_rows) + 1
    detail_headers = ['No. Soal', 'Tipe Soal', 'Bobot Maks', 'Jawaban Siswa',
                      'Kunci Jawaban', 'Status', 'Poin Didapat']
    for col_idx, h in enumerate(detail_headers, 1):
        ws.cell(row=detail_header_row, column=col_idx, value=h)
    _style_header_row(ws, detail_header_row, len(detail_headers))

    try:
        student_answers = json.loads(sub['answers_json']) if sub['answers_json'] else {}
    except Exception:
        student_answers = {}

    data_row = detail_header_row + 1
    total_earned = 0.0
    total_max = 0.0

    xlsx_evaluation = evaluate_answers_detailed(student_answers, questions)

    for q in questions:
        q_num = _normalize_q_num(q.get('number', ''))
        q_weight, earned = _build_question_detail_row(
            ws, data_row, q_num, q, student_answers, xlsx_evaluation
        )
        total_max += q_weight
        total_earned += earned
        data_row += 1

    if questions:
        ws.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=2)
        total_label = _style_data_cell(ws, data_row, 1, 'center')
        total_label.value = 'TOTAL'
        total_label.font = Font(bold=True, size=11, color=NAVY)

        total_max_cell = _style_data_cell(ws, data_row, 3, 'center')
        total_max_cell.value = round(total_max, 2)
        total_max_cell.font = Font(bold=True)

        total_earned_cell = _style_data_cell(ws, data_row, 7, 'center')
        total_earned_cell.value = round(total_earned, 2)
        total_earned_cell.font = Font(bold=True, size=12, color=NAVY)

        total_fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type='solid')
        for c in range(1, 8):
            _style_data_cell(ws, data_row, c).fill = total_fill
            _style_data_cell(ws, data_row, c).border = thin_border

    _set_column_widths(ws, [10, 16, 12, 24, 24, 18, 14])


def _generate_exam_xlsx(exam, submissions, questions, tz_offset=None):
    """Generate a professionally styled multi-sheet Excel workbook for a specific exam."""
    wb = openpyxl.Workbook()
    ident_fields = _parse_identity_fields(exam)

    _build_exam_summary_sheet(wb.active, exam, submissions, tz_offset)

    used_names = set()
    for sub in submissions:
        _build_student_detail_sheet(wb, sub, questions, tz_offset, used_names, ident_fields)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    safe_exam_name = ''.join(c for c in exam['name'] if c.isalnum() or c in ' _-').strip().replace(' ', '_')
    filename = f"Hasil_Ujian_{safe_exam_name}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"

    response = send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response
